"""Editable hydrology report for the same inputs used by DRENIA's rational method."""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.properties import CalcProperties

from flow import FLOW_PERIODS
from landuse import _YT_ROWS
from norm52 import NORM_USES


NAVY = '17324D'
TEAL = '147D83'
PALE = 'E6F3F3'
INPUT = 'E8F0FF'
GRAY = '607184'
WHITE = 'FFFFFF'


def _number(value):
    return float(str(value).replace(',', '.'))


def _heading(sheet, row, labels):
    for col, label in enumerate(labels, 1):
        cell = sheet.cell(row, col, label)
        cell.fill = PatternFill('solid', fgColor=NAVY)
        cell.font = Font(name='Aptos', size=10, bold=True, color=WHITE)
        cell.alignment = Alignment(wrap_text=True, vertical='center')
    sheet.row_dimensions[row].height = 33


def _title(sheet, text, last='J'):
    sheet.merge_cells(f'A1:{last}2')
    cell = sheet['A1']
    cell.value = text
    cell.font = Font(name='Aptos Display', size=19, bold=True, color=WHITE)
    cell.fill = PatternFill('solid', fgColor=NAVY)
    cell.alignment = Alignment(vertical='center', indent=1)
    sheet.row_dimensions[1].height = 27
    sheet.row_dimensions[2].height = 20


def export_excel(path: Path, *, basin: dict, landuse: dict, threshold: dict,
                 inputs: dict, channel: dict | None = None) -> Path:
    """Export a live Excel model; numbers and Ft come from the current app state."""
    path = Path(path)
    if threshold['p0_inicial_ponderado_mm'] is None:
        raise ValueError('Revisa todas las equivalencias de usos antes de exportar el Excel.')
    rows = threshold['filas']
    if not rows:
        raise ValueError('No hay superficies de usos para exportar.')
    wb = Workbook()
    wb.calculation = CalcProperties(calcMode='auto', fullCalcOnLoad=True, forceFullCalc=True)
    main = wb.active
    main.title = 'Caudales'
    uses = wb.create_sheet('Usos del suelo')
    yt = wb.create_sheet('Factores YT')

    # Surface classification and the area-weighted initial runoff threshold.
    _title(uses, 'Superficies y umbral de escorrentía', 'G')
    uses['A3'] = 'Áreas recortadas por la cuenca. Grupo de suelo: ' + threshold['grupo_suelo']
    uses['A4'] = 'Fuente: ' + str(landuse.get('fuente', 'GeoPackage de usos'))
    _heading(uses, 6, ['Uso SIOSE', 'Fila 5.2-IC', 'Categoría de la norma',
                       'Área (m²)', 'Área (ha)', 'P₀ⁱ (mm)', 'Área × P₀ⁱ'])
    for index, entry in enumerate(rows, 7):
        code = entry['fila_52ic']
        uses.cell(index, 1, entry['uso_siose'])
        uses.cell(index, 2, code)
        uses.cell(index, 3, NORM_USES[code].label if code else 'Sin equivalencia')
        uses.cell(index, 4, entry['area_m2'])
        uses.cell(index, 5, f'=D{index}/10000')
        uses.cell(index, 6, entry['p0_inicial_mm'])
        uses.cell(index, 7, f'=D{index}*F{index}')
        for col in (4, 5, 6, 7):
            uses.cell(index, col).number_format = '#,##0.000'
        if index % 2 == 0:
            for cell in uses[index][:7]:
                cell.fill = PatternFill('solid', fgColor='F2F7FA')
    total = 7 + len(rows)
    uses.cell(total, 3, 'TOTAL / P₀ⁱ PONDERADO')
    uses.cell(total, 4, f'=SUM(D7:D{total-1})')
    uses.cell(total, 5, f'=D{total}/10000')
    uses.cell(total, 6, f'=IF(D{total}>0,SUM(G7:G{total-1})/D{total},"")')
    uses.cell(total, 7, f'=SUM(G7:G{total-1})')
    for cell in uses[total][:7]:
        cell.fill = PatternFill('solid', fgColor=PALE)
        cell.font = Font(name='Aptos', bold=True, color=NAVY)
    uses.cell(total, 6).number_format = '0.000'
    uses.cell(total+2, 1, 'Superficie máscara MDT (km²)')
    uses.cell(total+2, 4, basin['area_km2'])
    uses.cell(total+3, 1, 'Diferencia vector / máscara (m²)')
    uses.cell(total+3, 4, f'=D{total}-D{total+2}*1000000')
    uses.cell(total+5, 1, 'P₀ⁱ procede de la tabla 2.3 de la 5.2-IC; pendiente <3 % y R/N donde corresponde.')
    for col, width in {'A': 34, 'B': 21, 'C': 48, 'D': 19, 'E': 16, 'F': 17, 'G': 20}.items():
        uses.column_dimensions[col].width = width
    uses.freeze_panes = 'D7'
    uses.auto_filter.ref = f'A6:G{total-1}'

    # Full YT lookup makes Cv an editable input in the workbook as well.
    _title(yt, 'Factores YT según Cv', 'H')
    yt['A3'] = 'Cv activo'
    yt['B3'] = '=Caudales!B14'
    yt['A4'] = 'Fila inferior'
    yt['B4'] = '=IF(OR(B3<0.3,B3>0.52),"",IFERROR(MATCH(B3,A9:A31,1),""))'
    yt['A5'] = 'Fila superior'
    yt['B5'] = '=IF(B4="","",MIN(23,B4+1))'
    yt['A6'] = 'Peso interpolación'
    yt['B6'] = '=IF(B4="","",IF(B4=B5,0,(B3-INDEX(A9:A31,B4))*100))'
    _heading(yt, 8, ['Cv', 2, 5, 10, 25, 50, 100, 500])
    for row, (cv_key, factors) in enumerate(sorted(_YT_ROWS.items()), 9):
        yt.cell(row, 1, cv_key/100).number_format = '0.00'
        for col, factor in enumerate((factors[i] for i in (0, 1, 2, 3, 4, 5, 7)), 2):
            yt.cell(row, col, factor).number_format = '0.000'
    yt['A33'] = 'Fuente: tabla de YT aportada para el estudio; interpolación lineal de Cv entre 0,30 y 0,52.'
    yt.column_dimensions['A'].width = 25
    for col in 'BCDEFGH':
        yt.column_dimensions[col].width = 15
    yt.freeze_panes = 'B9'

    # Calculation summary and editable drivers on the first sheet.
    _title(main, 'DRENIA · Caudales de cálculo')
    main['A3'] = 'Método racional 5.2-IC · cabecera: recorrido D8 más largo de la cuenca'
    for col, label, period in [('A', 'T5 · m³/s', 5), ('D', 'T100 · m³/s', 100),
                               ('G', 'T500 · m³/s', 500)]:
        main[f'{col}5'] = label
        main[f'{col}5'].font = Font(name='Aptos', bold=True, color=TEAL, size=12)
        result_row = 26 + FLOW_PERIODS.index(period)
        main[f'{col}6'] = f'=IF(ISNUMBER(I{result_row}),I{result_row},"Pendiente")'
        main[f'{col}6'].font = Font(name='Aptos Display', bold=True, size=17, color=NAVY)
        main[f'{col}6'].number_format = '0.0000'
    main['A8'] = 'Parámetros editables (azul)'
    main['A8'].font = Font(name='Aptos', bold=True, color=TEAL)
    labels = {9: 'Área cuenca (km²)', 10: 'Longitud recorrido (km)', 11: 'Pendiente media (m/m)',
              12: 'Tc manual (h; opcional)', 13: 'Tc utilizado (h)', 14: 'Cv',
              15: 'Precipitación media diaria (mm)', 16: 'βm regional', 17: 'I₁/Id',
              18: 'P₀ⁱ ponderado (mm)', 19: 'Ka', 20: 'Kt', 21: 'Grupo de suelo'}
    for row, label in labels.items():
        main.cell(row, 1, label)
    main['B9'] = basin['area_km2']
    main['B10'] = _number(inputs['longitud_km'])
    main['B11'] = _number(inputs['pendiente'])
    main['B12'] = _number(inputs['tc_h']) if inputs.get('tc_h') else None
    main['B13'] = '=IF(B12<>"",IF(B12>0,B12,""),IF(OR(B10<=0,B11<=0),"",0.3*(B10/B11^0.25)^0.76))'
    main['B14'] = _number(inputs['cv'])
    main['B15'] = _number(inputs['lluvia_media_mm'])
    main['B16'] = _number(inputs['beta_m'])
    main['B17'] = _number(inputs['i1_id'])
    main['B18'] = f"='Usos del suelo'!F{total}"
    main['B19'] = '=IF(B9<=0,"",IF(B9<1,1,1-LOG10(B9)/15))'
    main['B20'] = '=IF(B13="","",1+B13^1.25/(B13^1.25+14))'
    main['B21'] = threshold['grupo_suelo']
    for row in (9, 10, 11, 12, 14, 15, 16, 17):
        main.cell(row, 2).fill = PatternFill('solid', fgColor=INPUT)
    main['D9'] = 'Cabecera X'
    main['D10'] = 'Cabecera Y'
    main['D11'] = 'Cota cabecera (m)'
    main['D12'] = 'Cota vertido (m)'
    main['D13'] = 'Desnivel (m)'
    if channel:
        for row, key in ((9, 'cabecera_x'), (10, 'cabecera_y'), (11, 'cota_cabecera_m'),
                         (12, 'cota_vertido_m'), (13, 'desnivel_m')):
            main.cell(row, 5, channel[key])
    main['D15'] = ('Las cotas proceden del MDT original en los extremos del recorrido D8.'
                   if channel else 'Sin trazado automático: longitud y pendiente introducidas manualmente.')
    main['D16'] = 'Si editas Cv, YT se interpola en la hoja Factores YT.'
    main['D17'] = 'Ft de T10 y T50 se introduce tras interpolarlo fuera de DRENIA.'
    main['D18'] = 'Las celdas sin Ft quedan sin caudal.'
    main['D19'] = 'Los valores azules pueden modificarse; el resto se recalcula en Excel.'
    _heading(main, 25, ['T (años)', 'YT', 'Ft', 'Pd (mm)', 'P₀ (mm)', 'C',
                        'I (mm/h)', 'Kt', 'Q (m³/s)', 'Estado'])
    for row, period in enumerate(FLOW_PERIODS, 26):
        main.cell(row, 1, period)
        match = f'MATCH(A{row},\'Factores YT\'!$B$8:$H$8,0)'
        lo = f'INDEX(\'Factores YT\'!$B$9:$H$31,\'Factores YT\'!$B$4,{match})'
        hi = f'INDEX(\'Factores YT\'!$B$9:$H$31,\'Factores YT\'!$B$5,{match})'
        main.cell(row, 2, f'=IF(\'Factores YT\'!$B$4="","",{lo}*(1-\'Factores YT\'!$B$6)+{hi}*\'Factores YT\'!$B$6)')
        ft = inputs.get('ft', {}).get(str(period), '')
        main.cell(row, 3, _number(ft) if str(ft).strip() else None)
        main.cell(row, 4, f'=IF(OR(B{row}="",C{row}="",C{row}<=0,$B$15<=0),"",B{row}*$B$15)')
        main.cell(row, 5, f'=IF(OR(C{row}="",C{row}<=0,$B$18="",$B$16<=0),"",$B$18*$B$16*C{row})')
        ratio = f'(D{row}*$B$19/E{row})'
        main.cell(row, 6, f'=IF(OR(D{row}="",E{row}=""),"",IF(E{row}=0,1,IF(D{row}*$B$19<=E{row},0,(({ratio}-1)*({ratio}+23))/({ratio}+11)^2)))')
        main.cell(row, 7, f'=IF(OR(D{row}="",$B$13="",$B$17<=0,$B$19=""),"",D{row}*$B$19/24*$B$17^(3.5287-2.5287*$B$13^0.1))')
        main.cell(row, 8, '=IF($B$13="","",$B$20)')
        main.cell(row, 9, f'=IF(OR(F{row}="",G{row}="",H{row}=""),"",F{row}*G{row}*$B$9*H{row}/3.6)')
        main.cell(row, 10, f'=IF(OR(C{row}="",C{row}<=0),"Falta Ft",IF(I{row}="","Revisar entradas","Calculado"))')
        main.cell(row, 3).fill = PatternFill('solid', fgColor=INPUT)
        for col in range(2, 10):
            main.cell(row, col).number_format = '0.0000'
        if period in (5, 100, 500):
            for cell in main[row][:10]:
                if cell.coordinate != f'C{row}':
                    cell.fill = PatternFill('solid', fgColor=PALE)
                    cell.font = Font(name='Aptos', color=NAVY, bold=True)
    main['A35'] = 'Origen: máscara de cuenca y MDT de DRENIA; usos SIOSE recortados; tabla 2.3 de la 5.2-IC.'
    main['A36'] = 'El método del recorrido más largo maximiza longitud D8, no tiempo de viaje diferencial.'
    for col, width in {'A': 36, 'B': 18, 'C': 15, 'D': 31, 'E': 18,
                       'F': 16, 'G': 18, 'H': 15, 'I': 19, 'J': 20}.items():
        main.column_dimensions[col].width = width
    main.freeze_panes = 'D26'
    main.print_options.horizontalCentered = True
    main.sheet_properties.pageSetUpPr.fitToPage = True
    main.page_setup.fitToWidth = 1
    main.print_area = 'A1:J36'
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
