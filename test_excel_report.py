"""The export must preserve editable drivers, the source surfaces and missing Ft."""
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from excel_report import export_excel


class ExcelReportTest(unittest.TestCase):
    def test_export_summary_and_missing_period(self):
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp)/'DRENIA_caudales.xlsx'
            export_excel(output, basin={'area_km2': 23.87518103},
                landuse={'fuente': 'usos.gpkg'}, threshold={
                    'grupo_suelo': 'B', 'p0_inicial_ponderado_mm': 11,
                    'filas': [{'uso_siose': 'Casco', 'area_m2': 10000,
                               'fila_52ic': '11100', 'p0_inicial_mm': 1},
                              {'uso_siose': 'Cultivo', 'area_m2': 10000,
                               'fila_52ic': '21100', 'p0_inicial_mm': 21}]},
                inputs={'longitud_km': '10,21485', 'pendiente': '0,007861',
                        'tc_h': '', 'cv': '0,34', 'lluvia_media_mm': '37',
                        'beta_m': '2,15', 'i1_id': '10',
                        'ft': {'5': '0,88', '100': '1,38', '500': '1,62'}},
                channel={'cabecera_x': 400001, 'cabecera_y': 4400001,
                         'cota_cabecera_m': 699, 'cota_vertido_m': 618.7,
                         'desnivel_m': 80.3})
            wb = load_workbook(output)
            self.assertEqual(wb.sheetnames, ['Caudales', 'Usos del suelo', 'Factores YT'])
            sheet = wb['Caudales']
            self.assertEqual(sheet['B10'].value, 10.21485)
            self.assertEqual(sheet['B14'].value, .34)
            self.assertEqual(sheet['C27'].value, .88)
            self.assertIsNone(sheet['C26'].value)
            self.assertIn('Falta Ft', sheet['J26'].value)
            self.assertIn('I27', sheet['A6'].value)
            self.assertIn("'Usos del suelo'!F9", sheet['B18'].value)
            self.assertEqual(wb['Usos del suelo']['D7'].value, 10000)
            self.assertEqual(wb['Factores YT']['A13'].value, .34)
            self.assertIn('INDEX', sheet['B27'].value)
            self.assertEqual(wb.calculation.calcMode, 'auto')


if __name__ == '__main__':
    unittest.main()
